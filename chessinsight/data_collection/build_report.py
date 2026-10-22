"""수집 통계(data/stats)와 가공 데이터(data/processed)로 [데이터 수집 및 저장] 수집 데이터 보고서(.docx) 생성."""
import csv
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import HERE, PROCESSED, STATS

csv.field_size_limit(sys.maxsize if sys.maxsize < 2**31 else 2**31 - 1)

OUT = os.path.join(HERE, "..", "[데이터 수집 및 저장]수집 데이터_ChessInsight_AI.docx")
FONT = "맑은 고딕"
BROWN = "5B3A1E"
LABEL_FILL = "FDF3E8"
LINK = RGBColor(0x1F, 0x5F, 0xBF)
HEADER_TEXT = "ChessInsight AI: 대화형 체스 코칭 및 실력 분석·자동 리포팅 플랫폼"
BULLETS = {1: ("●", 0.6), 2: ("○", 1.4), 3: ("■", 2.2)}


def stats(name):
    with open(os.path.join(STATS, f"{name}.json"), encoding="utf-8") as f:
        return json.load(f)


def first_row(name):
    with open(os.path.join(PROCESSED, name), encoding="utf-8-sig") as f:
        return next(csv.DictReader(f))


def n(x):
    return f"{x:,}"


def short(text, limit=110):
    text = " ".join(str(text).split())
    return text if len(text) <= limit else text[:limit] + " …"


# ---------------------------------------------------------------- 서식 도우미

def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def para(doc_or_cell, text="", size=10, bold=False, align=None, space_after=4, color=None):
    p = doc_or_cell.add_paragraph()
    if text:
        set_font(p.add_run(text), size, bold, color)
    if align:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    return p


def bullet(doc, text, level=1):
    sym, indent = BULLETS[level]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(indent + 0.6)
    p.paragraph_format.first_line_indent = Cm(-0.6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(f"{sym}  "), 8 if level == 3 else 10)
    set_font(p.add_run(text), 10)
    return p


def bullets(doc, items, level=1):
    """items: 문자열 또는 (문자열, 하위 items) 튜플."""
    for item in items:
        if isinstance(item, tuple):
            bullet(doc, item[0], level)
            bullets(doc, item[1], level + 1)
        else:
            bullet(doc, item, level)


def heading(doc, text, size=13, before=14):
    p = para(doc, text, size=size, bold=True, space_after=8)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.keep_with_next = True
    return p


def shade(cell, fill):
    tc_pr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table, top=BROWN, top_sz=12, inside="000000", inside_sz=6, bottom=BROWN, bottom_sz=12):
    tbl_pr = table._element.tblPr
    el = OxmlElement("w:tblBorders")
    for edge, color, sz in (("top", top, top_sz), ("bottom", bottom, bottom_sz), ("insideH", inside, inside_sz)):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(sz))
        b.set(qn("w:color"), color)
        el.append(b)
    for edge in ("left", "right", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "nil")
        el.append(b)
    tbl_pr.append(el)


def cell_text(cell, text, size=9.5, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, link=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        if i:
            p = cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        set_font(r, size, bold, LINK if link else None)
        r.underline = link


def widths(table, cm_list):
    for row in table.rows:
        for cell, w in zip(row.cells, cm_list):
            cell.width = Cm(w)


def info_table(doc, rows):
    """양식의 데이터 정보 표: 좌측 라벨 칸(연한 배경) + 우측 내용. 내용은 문자열/줄 목록/(링크 텍스트,)."""
    t = doc.add_table(rows=len(rows), cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t, inside=BROWN, inside_sz=8)
    for (label, value), row in zip(rows, t.rows):
        shade(row.cells[0], LABEL_FILL)
        cell_text(row.cells[0], label, size=10, bold=True)
        if isinstance(value, tuple):
            cell_text(row.cells[1], value[0], size=10, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, link=True)
        else:
            cell_text(row.cells[1], value, size=10, align=WD_ALIGN_PARAGRAPH.LEFT)
        row.height = Cm(1.05)
    widths(t, [3.2, 12.8])
    para(doc, space_after=6)
    return t


def grid_table(doc, header, rows, col_cm, size=9):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders(t)
    for cell, h in zip(t.rows[0].cells, header):
        shade(cell, LABEL_FILL)
        cell_text(cell, h, size=size + 0.5, bold=True)
    for row, values in zip(t.rows[1:], rows):
        for cell, v in zip(row.cells, values):
            cell_text(cell, str(v), size=size)
    widths(t, col_cm)
    para(doc, space_after=6)
    return t


def field_table(doc, fields):
    """필드명·데이터 타입·설명·예시."""
    grid_table(doc, ["필드명", "데이터 타입", "설명", "예시"], fields, [3.4, 2.1, 4.3, 6.2])


def page_number(paragraph):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._element.append(el)
    set_font(run, 10)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- 본문

def cover(doc):
    for _ in range(7):
        para(doc)
    para(doc, "[데이터 수집 및 저장]", 24, True, WD_ALIGN_PARAGRAPH.CENTER, 36)
    para(doc, "수집 데이터 보고서", 24, True, WD_ALIGN_PARAGRAPH.CENTER, 10)
    para(doc, "SKN Family AI Camp 2기 : 최종 프로젝트 [  ]조", 14, True, WD_ALIGN_PARAGRAPH.CENTER, 0)
    for _ in range(5):
        para(doc)
    info_table(doc, [
        ("산출물 단계", "데이터 수집 및 저장"),
        ("평가 산출물", "수집 데이터 보고서"),
        ("제출 일자", "2026. 10. 22"),
        ("깃허브 경로", ("kty2001/final_mvp",)),
        ("작성 팀원", "[  ]"),
    ])
    page_break(doc)


def section_list(doc, d1, d2, d3, d4):
    heading(doc, "0. 수집 데이터 목록", 14, 0)
    bullets(doc, [
        ("목록 도출 기준", [
            "기획서(ChessInsight_AI_v2.md) 6장 데이터 수집 및 정책 고려 사항, 요구사항 정의서의 기능 ID 기준으로 필요 데이터 도출",
            "핵심 기능(엔진 분석 기반 수 설명, 자동 리포팅)에 직접 쓰이는 데이터를 P1으로 두고 우선 수집",
        ]),
    ])
    grid_table(doc, ["구분", "데이터", "출처 / 라이선스", "관련 기능 ID", "우선순위", "수집 상태"], [
        ["D1", "레이팅 구간별 대국 표본 (엔진 평가·남은 시간 포함)", "Lichess Open Database / CC0",
         "FAC-04, FAC-32, FAR-01", "P1 핵심", f"완료 ({n(d1['selected_total'])}판)"],
        ["D2", "테마 태그 퍼즐", "Lichess Puzzle DB / CC0", "FAC-14, FBF-26", "P1 핵심",
         f"완료 ({n(d2['usable_rows'])}개)"],
        ["D3", "오프닝 목록 (ECO·이름·수순)", "lichess-org/chess-openings / CC0", "FAC-06, FAC-22, FBF-29",
         "P1 핵심", f"완료 ({n(d3['usable_rows'])}개)"],
        ["D4", "사용자별 대국 이력 (가명화)", "Lichess API (공개 대국)", "FAC-26~31, FAC-34~36", "P1 핵심",
         f"완료 ({d4['users']}명, {n(d4['games'])}판)"],
        ["D5", "포지션 평가 캐시", "Lichess Evaluation DB / CC0", "FAC-09", "P2", "예정"],
        ["D6", "시나리오 평가셋", "D1·D4 2차 가공", "FAR-01, FAC-21", "P2", "예정"],
        ["D7", "한영 체스 용어집·사전 콘텐츠", "자체 작성·숙련자 검수", "FAD-01~06", "P3", "예정"],
        ["D8", "Chess.com 대국", "Published Data API", "FBF-10", "보류", "약관 검토 후 결정"],
    ], [1.1, 4.0, 3.6, 3.2, 1.6, 2.7], size=8.5)
    bullets(doc, [
        ("보류·후순위 사유", [
            "D5: 평가 캐시는 비용·지연 절감용으로, 핵심 기능은 Stockfish 직접 분석으로 동작함",
            "D6: D1·D4 대국에서 유형별 포지션을 추출하는 2차 가공 데이터로, 원천 수집 이후 진행",
            "D8: Chess.com API 약관의 \"자사 기능이나 IP 복제 금지\" 문구 검토 전까지 수집 보류(기획서 11장)",
        ]),
        ("저장 구조 (chessinsight/data_collection/data)", [
            "raw/: 원본 그대로 보존 (선별 대국 PGN, 퍼즐 원본 .csv.zst, 오프닝 원본 TSV)",
            "processed/: 검증·정제·가명화를 거친 CSV (UTF-8)",
            "stats/: 수집 건수·제외 사유 통계 JSON (본 보고서 수치의 출처)",
        ]),
    ])
    page_break(doc)


REJECT_KO = {
    "no_eval": "엔진 평가(%eval) 없음", "speed": "불릿·초불릿·통신 대국", "band_mismatch": "양측 레이팅 구간 불일치",
    "band_full": "해당 구간 목표 달성", "not_rated": "비레이팅 대국", "termination": "비정상 종료(기권 처리·규칙 위반 등)",
    "too_short": "20 ply 미만", "no_elo": "레이팅 없음", "no_clock": "남은 시간(%clk) 없음", "parse_error": "수순 파싱·합법성 오류",
    "duplicate": "중복 대국", "variant": "변형 체스", "variant_or_speed": "변형 체스·불릿 등 대상 외 속도",
    "user_under_min_games": "30판 미만 사용자(사용자 단위 제외)",
}


def section_games(doc, s):
    r = first_row("lichess_games.csv")
    heading(doc, "1. lichess_games.csv", 14, 0)
    info_table(doc, [
        ("데이터명", "레이팅 구간별 Lichess 대국 표본"),
        ("수집 대상", [f"① Lichess 월별 표준 대국 DB ({s['month']})",
                     "      ● 엔진 평가(%eval)·남은 시간(%clk) 주석이 있는 레이팅 대국",
                     "      ● 양측 레이팅이 같은 구간인 블리츠·래피드·클래식 대국"]),
        ("수집 목적", "레이팅대별 실수율 기준선, 실수 분류 기준표 보정, 시나리오 평가셋 원천 확보"),
        ("사용 예정 기능", ["실수 분류(FAC-04), 레이팅대 기준선 비교(FAC-32)", "시나리오 평가셋 구축(FAR-01)"]),
        ("수집 일자", "2026. 10. 22"),
        ("출처/저작권", ("Lichess Open Database (CC0)",)),
    ])
    heading(doc, "1.1. 수집 방법 및 자동화 절차", 12)
    bullets(doc, [
        ("수집 방식", [
            ("공개 DB 스트리밍", [
                "월별 압축 파일(.pgn.zst, 수십 GB)을 디스크에 저장하지 않고 HTTP 스트리밍으로 해제하며 필요한 대국만 선별함",
                f"구간별 목표 {n(s['per_band_target'])}판 달성 또는 압축 기준 3GB 읽기 시 자동 종료함",
            ]),
            ("헤더 기반 1차 필터", [
                "레이팅 대국, 표준 규칙, 블리츠·래피드·클래식, 정상 종료, 양측 같은 레이팅 구간, %eval·%clk 포함 여부를 헤더·수순 문자열로 판정함",
            ]),
            ("수순 검증", [
                "선별 대국만 MVP 서버의 PGN 파서(app/analysis.py parse_pgn)를 재사용해 합법성 검증·메타데이터 추출·해시 생성함",
            ]),
        ]),
        ("수집 도구", ["Python 3.12, requests, zstandard, python-chess",
                     "collect_lichess_games.py (월·구간별 목표·읽기 상한을 인자로 지정)"]),
        ("수집 주기", ["월 1회 (Lichess 월별 DB 공개 주기와 동일)"]),
    ])
    heading(doc, "1.2. 데이터 설명 및 구성", 12)
    bullets(doc, [
        f"스트리밍 검사 대국 수 : {n(s['streamed_games'])}판 (압축 {s['compressed_bytes_read'] / 1e6:,.0f} MB 읽음)",
        f"전체 수집 데이터 건수 : {n(s['selected_total'])}행",
        "사용할 데이터 건수 : 동일",
        "저장 포맷 : CSV (가명화), 원본 PGN (raw 보존)",
    ])
    grid_table(doc, ["레이팅 구간"] + list(s["selected_by_band"]),
               [["대국 수"] + [n(v) for v in s["selected_by_band"].values()]],
               [2.4] + [1.97] * len(s["selected_by_band"]), size=8.5)
    field_table(doc, [
        ["game_id", "object", "Lichess 대국 ID", r["game_id"]],
        ["rating_band", "object", "레이팅 구간 (7개)", r["rating_band"]],
        ["speed", "object", "시간 제어 구분", "blitz, rapid, classical"],
        ["white_id / black_id", "object", "가명화한 플레이어 ID", r["white_id"]],
        ["white_elo / black_elo", "int64", "대국 시점 레이팅", f"{r['white_elo']} / {r['black_elo']}"],
        ["result", "object", "대국 결과", r["result"]],
        ["termination", "object", "종료 사유", "Normal, Time forfeit"],
        ["played_at", "object", "대국 일자 (YYYY-MM-DD)", r["played_at"]],
        ["time_control", "object", "기본 시간+증가 시간(초)", r["time_control"]],
        ["eco / opening", "object", "ECO 코드 / 오프닝명", f"{r['eco']} / {r['opening']}"],
        ["n_plies", "int64", "총 반수(ply)", r["n_plies"]],
        ["has_eval / has_clock", "bool", "엔진 평가 / 남은 시간 주석 포함", "True / True"],
        ["pgn", "object", "가명화한 PGN 원문 (수별 %eval·%clk 포함)", short(r["pgn"].split("\n\n", 1)[-1])],
    ])
    heading(doc, "1.3. 데이터 품질 및 정합성 관리 방안", 12)
    rejected = s["rejected"]
    bullets(doc, [
        ("중복 제거 기준", [
            "Lichess 대국 ID(game_id)와 대국 해시(백·흑·날짜·시각·전체 수순의 SHA-256) 두 기준으로 중복 레코드 제거함",
            f"중복 제거 건수 : {n(rejected.get('duplicate', 0))}건",
        ]),
        ("정합성 검증 방법", [
            "필수 헤더(Event, WhiteElo, BlackElo, TimeControl, Result) 존재 여부와 레이팅 정수형 여부 검증함",
            "python-chess로 전 수순의 합법성 검증, 실패 대국은 제외함 "
            f"(제외 {n(rejected.get('parse_error', 0))}건)",
            "20 ply 미만 대국은 분석 가치가 낮아 제외함(기획서 6장 전처리 절차)",
        ]),
        ("필터 단계별 제외 현황", [f"{REJECT_KO.get(k, k)} : {n(v)}건" for k, v in rejected.items()]),
        ("표준화 전략", [
            "시간 제어 구분: Lichess 기준(기본 시간 + 40 × 증가 시간)으로 blitz·rapid·classical 산정",
            "레이팅 구간: ~1199부터 2200~까지 200점 단위 7개 구간으로 통일",
            "날짜: YYYY-MM-DD 형식으로 통일 (UTCDate 우선)",
            "개인정보 최소화: 사용자명은 솔트를 붙인 SHA-256 기반 가명 ID(u_ + 10자리)로 치환, 원본은 raw에만 보존",
        ]),
    ])
    page_break(doc)


def section_puzzles(doc, s):
    r = first_row("lichess_puzzles.csv")
    heading(doc, "2. lichess_puzzles.csv", 14, 0)
    info_table(doc, [
        ("데이터명", "Lichess 테마 태그 퍼즐"),
        ("수집 대상", ["① Lichess 퍼즐 DB 전체 (lichess_db_puzzle.csv.zst)",
                     "      ● 퍼즐 시작 포지션(FEN), 정답 수순, 퍼즐 레이팅, 테마 태그"]),
        ("수집 목적", "전술 모티프 검출기의 정답 라벨 확보, 약점 테마와 같은 퍼즐 추천"),
        ("사용 예정 기능", ["전술 모티프 검출 및 확정(FAC-14) 검증", "테마 퍼즐 연결(FBF-26), 내 실수 퀴즈 보조(FFJ-12)"]),
        ("수집 일자", "2026. 10. 22"),
        ("출처/저작권", ("Lichess Puzzle Database (CC0)",)),
    ])
    heading(doc, "2.1. 수집 방법 및 자동화 절차", 12)
    bullets(doc, [
        ("수집 방식", [
            ("공개 DB 다운로드", [
                f"월 단위로 갱신되는 퍼즐 DB 압축 파일({s['compressed_bytes'] / 1e6:.0f} MB)을 내려받아 raw에 원본 그대로 보존함",
                "압축 해제 스트림으로 한 줄씩 검증·정제하여 processed CSV로 저장함",
            ]),
            ("서비스 태그 매핑", [
                "Lichess 테마를 서비스 전술 태그 ID(요구사항 FAD-03·MVP 실수 태그 체계)로 변환하여 tactic_tags 컬럼으로 추가함",
            ]),
        ]),
        ("수집 도구", ["Python 3.12, requests, zstandard, python-chess", "collect_lichess_puzzles.py"]),
        ("수집 주기", ["월 1회 (Lichess 퍼즐 DB 갱신 주기와 동일)"]),
    ])
    heading(doc, "2.2. 데이터 설명 및 구성", 12)
    bullets(doc, [
        f"전체 수집 데이터 건수 : {n(s['total_rows'])}행",
        f"사용할 데이터 건수 : {n(s['usable_rows'])}행 (품질 기준: 레이팅 편차 ≤ 100, 풀이 횟수 ≥ 100)",
        "저장 포맷 : CSV, 원본 .csv.zst (raw 보존)",
    ])
    field_table(doc, [
        ["puzzle_id", "object", "퍼즐 ID", r["puzzle_id"]],
        ["fen", "object", "상대 수 직전 포지션", r["fen"]],
        ["moves", "object", "상대 수 + 정답 수순 (UCI)", r["moves"]],
        ["rating", "int64", "퍼즐 레이팅", r["rating"]],
        ["rating_deviation", "int64", "레이팅 편차", r["rating_deviation"]],
        ["popularity", "int64", "선호도 (-100~100)", r["popularity"]],
        ["nb_plays", "int64", "풀이 횟수", r["nb_plays"]],
        ["phase", "object", "국면 (opening·middlegame·endgame)", r["phase"] or "(태그 없음)"],
        ["tactic_tags", "object", "서비스 전술 태그 ID (공백 구분)", r["tactic_tags"] or "(해당 없음)"],
        ["themes", "object", "Lichess 원본 테마", r["themes"]],
        ["game_url", "object", "퍼즐이 나온 원본 대국", r["game_url"]],
        ["opening_tags", "object", "오프닝 태그", r["opening_tags"] or "(없음)"],
    ])
    tags = s["tactic_tag_counts"]
    grid_table(doc, ["서비스 전술 태그", "퍼즐 수", "서비스 전술 태그", "퍼즐 수"],
               [[a[0], n(a[1]), b[0] if b else "", n(b[1]) if b else ""]
                for a, b in zip(list(tags.items())[0::2], list(tags.items())[1::2] + [None])],
               [4.2, 3.9, 4.2, 3.9], size=8.5)
    heading(doc, "2.3. 데이터 품질 및 정합성 관리 방안", 12)
    inv = s["invalid"]
    bullets(doc, [
        ("정합성 검증 방법", [
            ("필수 필드 존재 여부 확인", [
                "PuzzleId, FEN, Moves, Rating, RatingDeviation, Popularity, NbPlays, Themes 8개 필드를 필수로 지정함",
                f"누락 레코드 {n(inv.get('missing_field', 0))}건, 숫자 타입 오류 {n(inv.get('bad_type', 0))}건 제외함",
            ]),
            ("수순 합법성 검증", [
                f"무작위 표본 {n(s['legality_sample'])}개(seed 고정)에 대해 FEN에서 정답 수순까지 모든 수의 합법성 검증함",
                f"검증 실패 {n(s['legality_sample_fail'])}건",
            ]),
        ]),
        ("중복 제거 기준", [f"PuzzleId 기준 중복 제거, 제거 건수 {n(inv.get('duplicate', 0))}건"]),
        ("품질 필터", [
            "레이팅 편차 100 초과(레이팅이 아직 안정되지 않은 퍼즐), 풀이 횟수 100 미만 퍼즐은 추천 대상에서 제외함",
        ]),
        ("표준화 전략", [
            f"Lichess 테마 {s['n_themes']}종 중 전술 모티프 {len(tags)}종을 서비스 태그 ID로 매핑 (예: attraction → decoy, discoveredAttack → discovered_attack)",
            "국면 테마(opening·middlegame·endgame)는 phase 컬럼으로 분리함",
            "과부하(overloading) 등 Lichess에 없는 서비스 태그는 퍼즐 라벨 없이 검출기 규칙으로만 다룸",
        ]),
    ])
    page_break(doc)


def section_openings(doc, s):
    r = first_row("openings.csv")
    heading(doc, "3. openings.csv", 14, 0)
    info_table(doc, [
        ("데이터명", "ECO 오프닝 목록"),
        ("수집 대상", ["① lichess-org/chess-openings 저장소 a~e.tsv",
                     "      ● ECO 코드, 오프닝·변형 이름, 대표 수순(PGN)"]),
        ("수집 목적", "오프닝 이론 수 식별, 해설에 쓰인 오프닝명 검증, 오프닝 사전 구축의 기준 목록"),
        ("사용 예정 기능", ["뻔한 수·이론 수 식별(FAC-06), 서버 측 출력 검증(FAC-22)",
                         "내 대국 오프닝 자동 식별(FBF-29), 오프닝 사전 구축(FAD-02)"]),
        ("수집 일자", "2026. 10. 22"),
        ("출처/저작권", ("lichess-org/chess-openings (CC0)",)),
    ])
    heading(doc, "3.1. 수집 방법 및 자동화 절차", 12)
    bullets(doc, [
        ("수집 방식", [
            ("GitHub 원본 파일 다운로드", ["ECO 그룹별 TSV 5개(a~e)를 내려받아 raw에 원본 그대로 보존함"]),
            ("수순 변환", [
                "python-chess로 PGN 수순을 재생하여 UCI 수순, 최종 포지션(EPD), ply 수 컬럼을 추가함",
                "EPD는 사용자 대국의 각 포지션과 대조하여 이론 수·이탈 지점을 찾는 키로 사용함",
            ]),
        ]),
        ("수집 도구", ["Python 3.12, requests, python-chess", "collect_openings.py"]),
        ("수집 주기", ["분기 1회 (저장소 갱신 시 수동 재실행)"]),
    ])
    heading(doc, "3.2. 데이터 설명 및 구성", 12)
    groups = ", ".join(f"{k} {n(v)}" for k, v in s["eco_groups"].items())
    bullets(doc, [
        f"전체 수집 데이터 건수 : {n(s['total_rows'])}행 (ECO 그룹별 {groups})",
        f"사용할 데이터 건수 : {n(s['usable_rows'])}행 (오프닝 계열 {s['families']}개)",
        "저장 포맷 : CSV, 원본 TSV (raw 보존)",
    ])
    field_table(doc, [
        ["eco", "object", "ECO 코드", r["eco"]],
        ["name", "object", "오프닝·변형 이름", r["name"]],
        ["family", "object", "오프닝 계열 (이름의 ':' 앞)", r["family"]],
        ["pgn", "object", "대표 수순 (SAN)", r["pgn"]],
        ["uci", "object", "대표 수순 (UCI)", r["uci"]],
        ["epd", "object", "수순 종료 포지션", r["epd"]],
        ["n_plies", "int64", "수순 길이(ply)", r["n_plies"]],
    ])
    heading(doc, "3.3. 데이터 품질 및 정합성 관리 방안", 12)
    inv = s["invalid"]
    bullets(doc, [
        ("정합성 검증 방법", [
            f"eco, name, pgn 필수 필드 존재 여부 검증 (누락 {n(inv.get('missing_field', 0))}건)",
            f"python-chess로 모든 수순의 합법성 검증 (오류 {n(inv.get('illegal_moves', 0))}건)",
        ]),
        ("중복 제거 기준", [f"[eco, name, pgn] 조합 기준 중복 제거 (제거 {n(inv.get('duplicate', 0))}건)",
                          f"같은 최종 포지션(EPD)을 공유하는 항목 {s['shared_epd_positions']}건으로 포지션 키의 유일성 확인"]),
        ("표준화 전략", [
            "수순 표기를 SAN·UCI 두 형식으로 함께 보관하여 PGN 파서·엔진 입력 양쪽에 사용함",
            "오프닝 계열(family) 컬럼으로 변형 단위와 계열 단위 집계를 모두 지원함 (오프닝별 성과 분석 FAC-28)",
            "오프닝 한글명은 한영 용어집(D7) 구축 시 같은 ECO·name 키로 연결 예정",
        ]),
    ])
    page_break(doc)


def section_users(doc, s):
    r = first_row("user_games.csv")
    u = first_row("user_profiles.csv")
    heading(doc, "4. user_games.csv / user_profiles.csv", 14, 0)
    info_table(doc, [
        ("데이터명", "사용자별 대국 이력 (개발·검증용 공개 대국)"),
        ("수집 대상", ["① Lichess API 사용자 대국 내보내기",
                     f"      ● D1 표본에 등장한 사용자 중 레이팅 구간별 {s['users_per_band_target']}명의 최근 대국",
                     "      ● 레이팅 블리츠·래피드·클래식, 남은 시간·엔진 평가·오프닝 포함"]),
        ("수집 목적", "한 사용자의 여러 대국에 걸친 반복 약점·기간 비교·10판 단위 추세 기능의 개발·검증"),
        ("사용 예정 기능", ["반복 약점 분석(FAC-26~31), 기간 지표 집계·10판 추세(FAC-34·35)",
                         "이전 리포트 대비 변화 판정(FAC-36)"]),
        ("수집 일자", "2026. 10. 22"),
        ("출처/저작권", ("Lichess API (공개 대국, Lichess 대국 데이터 CC0)",)),
    ])
    heading(doc, "4.1. 수집 방법 및 자동화 절차", 12)
    bullets(doc, [
        ("수집 방식", [
            ("API 호출", [
                f"GET /api/games/user/{{username}} 로 사용자당 최근 최대 {s['max_games_per_user']}판 요청 (rated, perfType=blitz·rapid·classical, clocks·evals·opening 포함)",
                "Lichess API 정책에 따라 한 번에 한 요청만 직렬로 보내고, 429 응답 시 60초 대기 후 재시도함",
            ]),
            ("사용자 선정", [
                "D1 표본에 등장한 사용자를 레이팅 구간별로 모아 무작위 순서(seed 고정)로 요청함",
                "기획서 6장 기준(1인당 30판 이상)을 충족하지 못한 사용자는 제외하고 다음 후보로 대체함",
            ]),
        ]),
        ("수집 도구", ["Python 3.12, requests, python-chess", "collect_user_games.py"]),
        ("수집 주기", ["일회성 (실제 서비스에서는 사용자 동의 후 연동 시점·리포트 주기마다 증분 수집, FBF-04·FBF-10)"]),
    ])
    heading(doc, "4.2. 데이터 설명 및 구성", 12)
    bullets(doc, [
        f"API로 받은 대국 수 : {n(s['fetched_games'])}판 (요청 {s['requests']}회)",
        f"전체 수집 데이터 건수 : 사용자 {s['users']}명, 대국 {n(s['games'])}행 "
        f"(1인당 {s['games_per_user_min']}~{s['games_per_user_max']}판, 엔진 평가 포함 {n(s['games_with_eval'])}판)",
        "사용할 데이터 건수 : 동일",
        "저장 포맷 : CSV 2종 (대국, 사용자 요약), 원본 PGN (raw 보존)",
    ])
    para(doc, "● user_games.csv", 10, True, space_after=4)
    field_table(doc, [
        ["user_id", "object", "가명화한 대상 사용자 ID", r["user_id"]],
        ["game_id", "object", "Lichess 대국 ID", r["game_id"]],
        ["user_color", "object", "대상 사용자의 색", r["user_color"]],
        ["user_elo / opponent_elo", "int64", "대국 시점 레이팅", f"{r['user_elo']} / {r['opponent_elo']}"],
        ["opponent_id", "object", "가명화한 상대 ID", r["opponent_id"]],
        ["result / user_score", "object / float64", "대국 결과 / 사용자 기준 점수", f"{r['result']} / {r['user_score']}"],
        ["termination", "object", "종료 사유", r["termination"]],
        ["played_at", "object", "대국 일자", r["played_at"]],
        ["speed / time_control", "object", "시간 제어", f"{r['speed']} / {r['time_control']}"],
        ["eco / opening", "object", "ECO 코드 / 오프닝명", f"{r['eco']} / {r['opening']}"],
        ["n_plies", "int64", "총 반수(ply)", r["n_plies"]],
        ["has_eval / has_clock", "bool", "엔진 평가 / 남은 시간 주석 포함", f"{r['has_eval']} / {r['has_clock']}"],
        ["pgn", "object", "가명화한 PGN 원문", short(r["pgn"].split("\n\n", 1)[-1])],
    ])
    para(doc, "● user_profiles.csv", 10, True, space_after=4)
    field_table(doc, [
        ["user_id", "object", "가명화한 사용자 ID", u["user_id"]],
        ["rating_band", "object", "선정 시 레이팅 구간", u["rating_band"]],
        ["n_games", "int64", "수집 대국 수", u["n_games"]],
        ["n_with_eval", "int64", "엔진 평가 포함 대국 수", u["n_with_eval"]],
        ["first_played / last_played", "object", "수집 대국 기간", f"{u['first_played']} ~ {u['last_played']}"],
        ["speeds", "object", "속도별 대국 수", u["speeds"]],
    ])
    heading(doc, "4.3. 데이터 품질 및 정합성 관리 방안", 12)
    rej = s["rejected"]
    bullets(doc, [
        ("정합성 검증 방법", [
            "D1과 같은 PGN 파서로 수순 합법성 검증, 20 ply 미만·비정상 종료·변형 체스 대국 제외함",
            ("제외 현황", [f"{REJECT_KO.get(k, k)} : {n(v)}건" for k, v in rej.items()] or ["없음"]),
        ]),
        ("중복 제거 기준", ["사용자별 game_id 기준 중복 제거, 같은 사용자를 두 구간에서 중복 선정하지 않음"]),
        ("표준화 전략", [
            "D1과 같은 컬럼명·속도 구분·날짜 형식·가명화 규칙을 사용하여 두 데이터셋을 같은 분석 코드로 처리함",
            "사용자 관점 컬럼(user_color, user_score, opponent_elo)을 추가하여 리포트 집계(승률·상대 수준 맥락 FAC-37)에 바로 사용함",
        ]),
        ("개인정보·정책 준수", [
            "공개 대국만 수집하고 사용자명은 가명 ID로 치환, 원본 사용자명은 로컬 raw 파일에만 보관하며 저장소에 커밋하지 않음",
            "실제 서비스에서는 명시적 동의를 받은 사용자 대국만 수집하고 삭제 요청 시 원본·분석 결과 삭제함(기획서 6장)",
        ]),
    ])


def main():
    d1, d2, d3, d4 = stats("lichess_games"), stats("lichess_puzzles"), stats("openings"), stats("user_games")
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)

    hp = sec.header.paragraphs[0]
    set_font(hp.add_run(HEADER_TEXT), 9)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    page_number(fp)

    cover(doc)
    section_list(doc, d1, d2, d3, d4)
    section_games(doc, d1)
    section_puzzles(doc, d2)
    section_openings(doc, d3)
    section_users(doc, d4)
    doc.save(OUT)
    print("saved:", os.path.abspath(OUT))


if __name__ == "__main__":
    main()
